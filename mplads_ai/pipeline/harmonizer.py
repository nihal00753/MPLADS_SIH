"""
Raw file ingestion and schema harmonization for MPLADS data.

Reads data from multiple file formats (CSV, TSV, XLSX, JSON) and maps
inconsistent government portal column names to canonical field names
expected by the MPLADS AI/ML pipeline.

Dependencies: ``openpyxl`` (for ``.xlsx`` support), stdlib otherwise.
"""

from __future__ import annotations

import csv
import io
import json
import logging
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Union

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Canonical Column Alias Map
# ---------------------------------------------------------------------------
# Keys are the canonical column names used in the MPLADS pipeline.
# Values are lists of known aliases from various government portals
# (eSAKSHI, PFMS, State portals, manually entered spreadsheets).

CANONICAL_COLUMN_MAP: Dict[str, List[str]] = {
    "unique_work_number": [
        "Work_Code", "Work_ID", "UWN", "work_code", "work_id",
        "Unique Work Number", "Unique_Work_Number", "WorkCode",
        "work_number", "Work Number", "Work_No", "work_no",
    ],
    "state": [
        "State", "STATE", "state_name", "State_Name", "StateName",
    ],
    "nodal_district": [
        "District", "Nodal_District", "district", "nodal_district_name",
        "Nodal District", "NodalDistrict",
    ],
    "implementing_district": [
        "Implementing_District", "impl_district", "Implementing District",
        "ImplementingDistrict", "Impl_District",
    ],
    "house_name": [
        "House_Name", "House", "house", "HouseName",
    ],
    "member_type": [
        "Member_Type", "MemberType", "member_type_name",
    ],
    "mp_name": [
        "MP_Name", "Mp_Name", "member_name", "MP Name", "MpName",
        "Member Name", "Member_Name", "mp",
    ],
    "constituency": [
        "Constituency", "constituency_name", "Constituency_Name",
        "ConstituencyName",
    ],
    "work_category": [
        "Category", "Work_Category", "work_type", "Work_Type",
        "WorkCategory", "Work Category", "WorkType",
    ],
    "work_name": [
        "Work_Name", "work_description", "Work_Description", "Title",
        "WorkName", "Work Name", "Description",
    ],
    "sanction_amount": [
        "Sanction_Amt", "Amount_Rs", "Approved_Cost", "amount", "Amount",
        "Sanctioned_Amount", "SanctionAmount", "Sanction Amount",
        "Cost", "sanctioned_cost", "Sanctioned_Cost",
    ],
    "date_of_receipt_of_work_proposal_from_mp": [
        "Proposal_Date", "proposal_date", "Date_Proposal",
        "ProposalDate", "Proposal Date", "Date of Proposal",
        "Receipt_Date", "receipt_date",
    ],
    "date_of_administrative_approval": [
        "Approval_Date", "approval_date", "Date_Approval",
        "Admin_Approval_Date", "ApprovalDate", "Approval Date",
        "Date of Approval", "AdminApprovalDate",
    ],
    "implementing_agency_name": [
        "Agency", "Implementing_Agency", "agency_name", "Agency_Name",
        "AgencyName", "Implementing Agency", "ImplementingAgency",
    ],
    "vendor_id": [
        "Vendor_ID", "vendor_code", "Vendor_Code", "contractor_id",
        "VendorID", "Vendor ID", "ContractorID", "Contractor_ID",
    ],
    "vendor_name": [
        "Vendor_Name", "vendor", "Contractor_Name", "contractor_name",
        "VendorName", "Vendor Name", "Contractor", "ContractorName",
    ],
    "work_status": [
        "Status", "Work_Status", "work_state", "Project_Status",
        "WorkStatus", "Work Status", "ProjectStatus",
    ],
    "physical_progress_pct": [
        "Physical_Progress", "physical_pct", "Physical_Pct",
        "PhysicalProgress", "Physical Progress", "phys_pct",
    ],
    "financial_progress_pct": [
        "Financial_Progress", "financial_pct", "Financial_Pct",
        "FinancialProgress", "Financial Progress", "fin_pct",
    ],
    "unit": [
        "Unit", "unit_type", "UnitType",
    ],
    "note": [
        "Note", "Notes", "Remarks", "remarks", "Comment", "comment",
    ],
    # Payment-specific columns
    "payment_id": [
        "Payment_ID", "PaymentID", "payment_code", "PaymentCode",
    ],
    "payment_date": [
        "Payment_Date", "PaymentDate", "payment_dt", "Date_Payment",
    ],
    "payment_amount": [
        "Payment_Amount", "PaymentAmount", "paid_amount", "PaidAmount",
        "Payment Amount", "amount_paid",
    ],
    # Vendor master columns
    "pan": ["PAN", "pan_number", "PAN_Number", "Pan"],
    "gstin": ["GSTIN", "gstin_number", "GSTIN_Number", "GST", "gst"],
    "is_hub": ["IsHub", "is_hub_vendor", "Hub", "hub"],
}


def _build_reverse_lookup() -> Dict[str, str]:
    """Build a case-insensitive alias → canonical_name lookup table."""
    lookup: Dict[str, str] = {}
    for canonical, aliases in CANONICAL_COLUMN_MAP.items():
        # The canonical name itself should map to itself
        lookup[canonical.lower().strip()] = canonical
        for alias in aliases:
            lookup[alias.lower().strip()] = canonical
    return lookup


_REVERSE_LOOKUP = _build_reverse_lookup()


# ---------------------------------------------------------------------------
# Multi-Format File Reader
# ---------------------------------------------------------------------------

def read_raw_file(
    file_path_or_buffer: Union[str, Path, BinaryIO, io.StringIO],
    file_type: str = "auto",
) -> List[Dict[str, str]]:
    """
    Read raw data from CSV, TSV, XLSX, or JSON into a list of dicts.

    Args:
        file_path_or_buffer: File path (str/Path) or file-like object
            (BytesIO/StringIO).
        file_type: One of ``"auto"``, ``"csv"``, ``"tsv"``, ``"xlsx"``,
            ``"json"``.  ``"auto"`` detects from file extension.

    Returns:
        List of row dicts with original (un-harmonized) column names.

    Raises:
        ValueError: If file type cannot be determined or is unsupported.
        FileNotFoundError: If a path is given but the file does not exist.
    """
    # Determine file type
    if file_type == "auto":
        if isinstance(file_path_or_buffer, (str, Path)):
            ext = Path(file_path_or_buffer).suffix.lower()
            type_map = {".csv": "csv", ".tsv": "tsv", ".xlsx": "xlsx", ".json": "json"}
            file_type = type_map.get(ext, "")
            if not file_type:
                raise ValueError(
                    f"Cannot auto-detect file type from extension '{ext}'. "
                    f"Supported: .csv, .tsv, .xlsx, .json"
                )
        else:
            raise ValueError(
                "Cannot auto-detect file type from a buffer. "
                "Please specify file_type explicitly."
            )

    file_type = file_type.lower().strip()

    if file_type == "csv":
        return _read_csv(file_path_or_buffer, delimiter=",")
    elif file_type == "tsv":
        return _read_csv(file_path_or_buffer, delimiter="\t")
    elif file_type == "xlsx":
        return _read_xlsx(file_path_or_buffer)
    elif file_type == "json":
        return _read_json(file_path_or_buffer)
    else:
        raise ValueError(f"Unsupported file type: '{file_type}'")


def _read_csv(
    source: Union[str, Path, BinaryIO, io.StringIO],
    delimiter: str = ",",
) -> List[Dict[str, str]]:
    """Read CSV/TSV from path or buffer."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        with open(path, encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            return list(reader)
    elif isinstance(source, io.StringIO):
        reader = csv.DictReader(source, delimiter=delimiter)
        return list(reader)
    else:
        # BinaryIO — decode to string first
        text = source.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8")
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        return list(reader)


def _read_xlsx(source: Union[str, Path, BinaryIO]) -> List[Dict[str, str]]:
    """Read first sheet of an XLSX file into list of dicts."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required for .xlsx support. "
            "Install it with: pip install openpyxl>=3.1"
        )

    if isinstance(source, (str, Path)):
        wb = openpyxl.load_workbook(str(source), read_only=True, data_only=True)
    else:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)

    ws = wb.active
    if ws is None:
        wb.close()
        return []

    rows_iter = ws.iter_rows(values_only=True)

    # First row = headers
    try:
        headers_raw = next(rows_iter)
    except StopIteration:
        wb.close()
        return []

    headers = [str(h).strip() if h is not None else f"_col_{i}" for i, h in enumerate(headers_raw)]

    result: List[Dict[str, str]] = []
    for row_values in rows_iter:
        row_dict: Dict[str, str] = {}
        for i, val in enumerate(row_values):
            if i < len(headers):
                row_dict[headers[i]] = str(val) if val is not None else ""
        if any(v.strip() for v in row_dict.values()):  # skip fully blank rows
            result.append(row_dict)

    wb.close()
    return result


def _read_json(source: Union[str, Path, BinaryIO, io.StringIO]) -> List[Dict[str, str]]:
    """Read JSON (expects a list of dicts)."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    elif isinstance(source, io.StringIO):
        data = json.load(source)
    else:
        text = source.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8")
        data = json.loads(text)

    if not isinstance(data, list):
        raise ValueError("JSON root must be a list of objects (dicts).")

    # Convert all values to strings for consistency with CSV reader
    return [
        {str(k): str(v) if v is not None else "" for k, v in row.items()}
        for row in data
        if isinstance(row, dict)
    ]


# ---------------------------------------------------------------------------
# Column Harmonizer
# ---------------------------------------------------------------------------

def harmonize_columns(
    rows: List[Dict[str, str]],
    column_map: Optional[Dict[str, List[str]]] = None,
) -> List[Dict[str, str]]:
    """
    Map incoming column names to canonical MPLADS field names.

    Uses the global ``CANONICAL_COLUMN_MAP`` unless a custom map is provided.
    Unrecognized columns are passed through unchanged (no data loss).

    Args:
        rows: List of row dicts with original column names.
        column_map: Optional override for the alias map.

    Returns:
        New list of row dicts with canonical column names.
    """
    if not rows:
        return []

    # Build lookup from the provided or default map
    if column_map is not None:
        lookup: Dict[str, str] = {}
        for canonical, aliases in column_map.items():
            lookup[canonical.lower().strip()] = canonical
            for alias in aliases:
                lookup[alias.lower().strip()] = canonical
    else:
        lookup = _REVERSE_LOOKUP

    # Determine the column name mapping from the first row's keys
    sample_keys = list(rows[0].keys())
    key_remap: Dict[str, str] = {}
    for original_key in sample_keys:
        canonical = lookup.get(original_key.lower().strip())
        key_remap[original_key] = canonical if canonical else original_key

    mapped_count = sum(1 for o, c in key_remap.items() if o != c)
    if mapped_count > 0:
        logger.info(
            "Harmonized %d/%d columns to canonical names",
            mapped_count, len(key_remap),
        )

    # Apply the mapping to all rows
    result: List[Dict[str, str]] = []
    for row in rows:
        new_row: Dict[str, str] = {}
        for orig_key, value in row.items():
            new_key = key_remap.get(orig_key, orig_key)
            new_row[new_key] = value
        result.append(new_row)

    return result
